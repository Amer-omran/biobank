import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from sra import connectors, demo_data, pipeline, reports
from sra.analytics import Filters
from sra.hf_dataset import profile_dataframe


@pytest.fixture
def report_data(conn):
    demo_data.load_demo_into_db(conn)
    pipeline.analyze_reviews(conn)
    return reports.build_report_data(conn, Filters(restaurants=["بيت المندي (تجريبي)"]))


def test_report_data(report_data):
    assert report_data["n_reviews"] == 130
    assert report_data["caveats"][0].startswith("⚠")
    assert any("ليس عدد مرات الشراء" in c for c in report_data["caveats"])
    assert report_data["recommendations"]
    assert any("من 130" in line for line in reports.summary_lines(report_data))


def test_excel_export_rtl(report_data):
    wb = load_workbook(io.BytesIO(reports.export_excel(report_data)))
    assert {"الملخص", "توزيع المشاعر", "الجوانب", "الوجبات", "التعليقات"} <= set(wb.sheetnames)
    assert all(ws.sheet_view.rightToLeft for ws in wb.worksheets)
    assert wb["التعليقات"].max_row == 131


def test_pdf_export(report_data):
    pdf = reports.export_pdf(report_data)
    assert pdf.startswith(b"%PDF") and len(pdf) > 5000
    assert b"Amiri" in pdf


def test_empty_report(conn):
    data = reports.build_report_data(conn)
    assert data["n_reviews"] == 0
    assert reports.export_pdf(data).startswith(b"%PDF")
    assert reports.summary_lines(data) == ["لا توجد تعليقات محللة تطابق الاختيار."]


def test_shape_arabic_wraps_in_logical_order():
    reports._register_fonts()
    lines = reports.wrap_rtl("كلمة " * 40, "Amiri", 11, 100)
    assert len(lines) > 3


def test_x_connector_requires_token(monkeypatch):
    monkeypatch.delenv("X_BEARER_TOKEN", raising=False)
    with pytest.raises(connectors.ConnectorError):
        connectors.XRecentSearchConnector("مطعم")


def test_x_connector_parses_and_hides_authors():
    captured = {}

    def fake_get(url, headers):
        captured["url"], captured["headers"] = url, headers
        return 200, {"data": [{"id": "1", "text": "الأكل لذيذ @someone", "created_at": "2025-03-01T10:00:00Z",
                               "author_id": "999"}]}

    conn = connectors.XRecentSearchConnector("مطعم", 5, "مطعم أ", bearer_token="secret", http_get=fake_get)
    df = conn.fetch()
    assert list(df.columns) == ["text", "review_date", "source", "restaurant_name"]
    assert df.iloc[0]["review_date"] == "2025-03-01" and "@someone" not in df.iloc[0]["text"]
    assert "max_results=10" in captured["url"] and "lang%3Aar" in captured["url"]
    assert captured["headers"]["Authorization"] == "Bearer secret"


@pytest.mark.parametrize("status", [401, 403, 429, 500])
def test_x_connector_errors(status):
    conn = connectors.XRecentSearchConnector("q", bearer_token="t", http_get=lambda u, h: (status, {}))
    with pytest.raises(connectors.ConnectorError):
        conn.fetch()


def test_profile_dataframe_handles_nested_columns():
    df = pd.DataFrame({
        "sentence": ["الأكل لذيذ", "الخدمة سيئة", "مكان جميل"],
        "triplets": [[{"a": "الأكل", "o": "لذيذ", "s": "POS"}], [], [{"a": "مكان"}]],
        "label": ["POS", "NEG", "POS"],
    })
    profile = profile_dataframe(df)
    cols = {c["name"]: c for c in profile["columns"]}
    assert profile["n_rows"] == 3
    assert cols["triplets"]["python_type"] == "list"
    assert cols["label"]["value_counts"] == {"POS": 2, "NEG": 1}
    assert cols["sentence"]["arabic_share"] == 1.0
