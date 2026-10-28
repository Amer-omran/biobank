"""إنشاء التقارير: ملخص إحصائي + تصدير Excel وPDF عربي.

PDF العربي يحتاج أمرين لا توفرهما ReportLab تلقائيًا:
1. تشكيل الحروف (وصلها) عبر arabic_reshaper.
2. ترتيب النص من اليمين لليسار عبر python-bidi.
ونقسم الأسطر بأنفسنا قبل عكس الاتجاه، وإلا ظهرت أسطر الفقرة الطويلة بترتيب معكوس.
الخط المستخدم Amiri (رخصة SIL OFL، مرفق في assets/fonts).
"""

from __future__ import annotations

import io
from datetime import datetime
from xml.sax.saxutils import escape

import pandas as pd

from . import analytics, config
from .analytics import Filters, fraction_text
from .text_processing import EMOJI_RE

RECOMMENDATIONS = {
    "food_quality": "مراجعة جودة الأطباق وثبات الطعم ودرجة الحرارة عند التقديم.",
    "price": "مراجعة التسعير أو تقديم عروض وقيمة أوضح مقابل السعر.",
    "quantity": "مراجعة أحجام الحصص مقارنة بالسعر وتوقعات العملاء.",
    "service": "تدريب فريق الخدمة وتقليل وقت الانتظار.",
    "delivery": "مراجعة زمن التوصيل والتغليف الحافظ للحرارة مع شركاء التوصيل.",
    "cleanliness": "تشديد إجراءات النظافة في الصالة ودورات المياه والمطبخ.",
    "ambiance": "تحسين الجلسات وتقليل الازدحام والضوضاء.",
}
MIN_FOR_RECOMMENDATION = 10
RECOMMENDATION_NEG_PCT = 35.0


# =================================================================== البيانات
def build_report_data(conn, filters: Filters | None = None, title: str | None = None) -> dict:
    filters = filters or Filters()
    reviews = analytics.analyzed_reviews(conn, filters)
    aspects = analytics.aspects_frame(conn, filters)
    dishes = analytics.dishes_frame(conn, filters)
    n = len(reviews)

    aspect_table = analytics.aspect_summary(aspects)
    dish_table = analytics.dish_summary(dishes).head(15)
    trend = analytics.sentiment_trend(reviews, "M")
    restaurants = filters.restaurants or sorted(reviews["restaurant_name"].dropna().unique().tolist())

    caveats = [
        "نتائج التحليل آلية وتقديرية، وقد تخطئ في السخرية والعامية غير المألوفة.",
        "عدد مرات ذكر الوجبة مؤشر تفضيل مستنتج من التعليقات، وليس عدد مرات الشراء أو المبيعات.",
        "درجة الثقة تقديرية وليست احتمالًا معايرًا إحصائيًا.",
    ]
    if n and reviews["is_synthetic"].astype(bool).any():
        caveats.insert(0, "⚠ يحتوي التقرير على بيانات تجريبية اصطناعية؛ النتائج ليست نتائج بحثية فعلية.")
    if 0 < n < config.MIN_REVIEWS_FOR_RANKING:
        caveats.append(f"عدد التعليقات ({n}) أقل من {config.MIN_REVIEWS_FOR_RANKING}؛ النتائج غير كافية للتعميم.")

    recommendations = []
    for row in aspect_table.to_dict("records"):
        if row["n"] >= MIN_FOR_RECOMMENDATION and row["negative_pct"] >= RECOMMENDATION_NEG_PCT:
            recommendations.append(
                f"{row['aspect_ar']}: {fraction_text(int(row['negative']), int(row['n']))} من الإشارات سلبية. "
                + RECOMMENDATIONS[row["aspect"]]
            )
    if not recommendations and n:
        recommendations.append("لا يوجد جانب تتجاوز شكاواه الحد المعتمد للتوصية "
                               f"({RECOMMENDATION_NEG_PCT:.0f}% مع {MIN_FOR_RECOMMENDATION} إشارات على الأقل).")

    opinions = aspects[aspects["has_opinion"] == 1] if not aspects.empty else aspects
    complaints = (opinions[opinions["sentiment"] == "negative"].sort_values("score")
                  .drop_duplicates("evidence").head(8)) if not opinions.empty else opinions
    praises = (opinions[opinions["sentiment"] == "positive"].sort_values("score", ascending=False)
               .drop_duplicates("evidence").head(8)) if not opinions.empty else opinions

    period_from = filters.start_date or (reviews["review_date"].dropna().min() if n else None)
    period_to = filters.end_date or (reviews["review_date"].dropna().max() if n else None)
    ratings = reviews["original_rating"].dropna() if n else pd.Series(dtype=float)
    return {
        "title": title or "تقرير تحليل مشاعر العملاء",
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "restaurants": restaurants,
        "period": (period_from, period_to),
        "datasets": sorted(reviews["dataset"].unique().tolist()) if n else [],
        "model_versions": sorted(reviews["model_version"].unique().tolist()) if n else [],
        "n_reviews": n,
        "n_mixed": int(reviews["is_mixed"].sum()) if n else 0,
        "n_uncertain": int(reviews["is_uncertain"].sum()) if n else 0,
        "avg_rating": round(float(ratings.mean()), 2) if len(ratings) else None,
        "n_rated": int(len(ratings)),
        "distribution": analytics.sentiment_distribution(reviews),
        "aspects": aspect_table,
        "dishes": dish_table,
        "trend": trend,
        "complaints": complaints,
        "praises": praises,
        "recommendations": recommendations,
        "caveats": caveats,
        "reviews": reviews,
    }


def summary_lines(data: dict) -> list[str]:
    """ملخص نصي قصير يُستخدم في الواجهة وفي PDF."""
    if not data["n_reviews"]:
        return ["لا توجد تعليقات محللة تطابق الاختيار."]
    p_from, p_to = data["period"]
    lines = [
        f"المطاعم: {'، '.join(data['restaurants']) if data['restaurants'] else 'غير محدد'}",
        f"الفترة: {p_from or 'غير محدد'} إلى {p_to or 'غير محدد'}",
        f"عدد التعليقات المحللة: {data['n_reviews']}",
    ]
    for row in data["distribution"].to_dict("records"):
        lines.append(f"{row['label']}: {row['text']}")
    lines.append(f"تعليقات مختلطة المشاعر: {fraction_text(data['n_mixed'], data['n_reviews'])}")
    lines.append(f"نتائج غير مؤكدة (ثقة منخفضة): {fraction_text(data['n_uncertain'], data['n_reviews'])}")
    if data["avg_rating"] is not None:
        lines.append(f"متوسط التقييم الأصلي: {data['avg_rating']} (من {data['n_rated']} تقييمًا)")
    return lines


# ===================================================================== Excel
def export_excel(data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F5F4A")
    right = Alignment(horizontal="right", vertical="top", wrap_text=True)

    def add_sheet(title, df: pd.DataFrame, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.sheet_view.rightToLeft = True
        ws.append(list(df.columns))
        for cell in ws[1]:
            cell.font, cell.fill, cell.alignment = header_font, header_fill, right
        for row in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
        for col_cells in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = min(60, max(10, width + 2))
            for c in col_cells[1:]:
                c.alignment = right
        ws.freeze_panes = "A2"
        return ws

    summary = pd.DataFrame({"البند": ["العنوان", "تاريخ الإنشاء", "مجموعات البيانات", "إصدارات النموذج"]
                            + [line.split(":")[0] for line in summary_lines(data)],
                            "القيمة": [data["title"], data["generated_at"], "، ".join(data["datasets"]),
                                       "، ".join(data["model_versions"])]
                            + [":".join(line.split(":")[1:]).strip() for line in summary_lines(data)]})
    add_sheet("الملخص", summary, first=True)

    dist = data["distribution"].rename(columns={"label": "الشعور", "count": "العدد", "total": "المقام",
                                                "pct": "النسبة %"})[["الشعور", "العدد", "المقام", "النسبة %"]]
    add_sheet("توزيع المشاعر", dist)

    if not data["aspects"].empty:
        add_sheet("الجوانب", data["aspects"].rename(columns={
            "aspect_ar": "الجانب", "n": "عدد الإشارات", "positive": "إيجابي", "negative": "سلبي",
            "neutral": "محايد", "positive_pct": "إيجابي %", "negative_pct": "سلبي %", "neutral_pct": "محايد %",
        })[["الجانب", "عدد الإشارات", "إيجابي", "سلبي", "محايد", "إيجابي %", "سلبي %", "محايد %"]])
    if not data["dishes"].empty:
        add_sheet("الوجبات", data["dishes"].rename(columns={
            "dish": "الوجبة", "category": "الفئة", "mentions": "تعليقات ذكرتها", "positive": "إيجابي",
            "negative": "سلبي", "neutral": "محايد", "positive_pct": "إيجابي %", "negative_pct": "سلبي %",
        })[["الوجبة", "الفئة", "تعليقات ذكرتها", "إيجابي", "سلبي", "محايد", "إيجابي %", "سلبي %"]])
    if not data["trend"].empty:
        add_sheet("الاتجاه الزمني", data["trend"].rename(columns={
            "period": "الشهر", "n": "العدد", "positive": "إيجابي", "negative": "سلبي", "neutral": "محايد",
        })[["الشهر", "العدد", "إيجابي", "سلبي", "محايد"]])
    add_sheet("التوصيات والتنبيهات", pd.DataFrame(
        {"النوع": ["توصية"] * len(data["recommendations"]) + ["تنبيه"] * len(data["caveats"]),
         "النص": data["recommendations"] + data["caveats"]}))
    if data["n_reviews"]:
        reviews = data["reviews"][["text", "restaurant_name", "city", "review_date", "original_rating",
                                   "sentiment", "confidence", "is_mixed", "is_uncertain", "cues"]].copy()
        reviews["sentiment"] = reviews["sentiment"].map(config.SENTIMENT_AR)
        reviews["is_mixed"] = reviews["is_mixed"].map({1: "نعم", 0: "لا"})
        reviews["is_uncertain"] = reviews["is_uncertain"].map({1: "نعم", 0: "لا"})
        reviews.columns = ["التعليق", "المطعم", "المدينة", "التاريخ", "التقييم الأصلي", "الشعور",
                           "درجة الثقة", "مختلط", "غير مؤكد", "كلمات الرأي"]
        add_sheet("التعليقات", reviews)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ======================================================================= PDF
_FONTS_REGISTERED = False


def _register_fonts() -> None:
    global _FONTS_REGISTERED
    if _FONTS_REGISTERED:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if not config.ARABIC_FONT_REGULAR.exists():
        raise FileNotFoundError(f"الخط العربي غير موجود: {config.ARABIC_FONT_REGULAR}")
    pdfmetrics.registerFont(TTFont("Amiri", str(config.ARABIC_FONT_REGULAR)))
    pdfmetrics.registerFont(TTFont("Amiri-Bold", str(config.ARABIC_FONT_BOLD)))
    _FONTS_REGISTERED = True


def shape_arabic(text: str) -> str:
    """تشكيل الحروف وعكس الاتجاه لسطر واحد."""
    import arabic_reshaper
    from bidi.algorithm import get_display

    # خط Amiri لا يحتوي الرموز التعبيرية، فنحذفها من PDF فقط
    text = EMOJI_RE.sub("", str(text)).replace("\u2764", "").strip()
    return get_display(arabic_reshaper.reshape(text))


def wrap_rtl(text: str, font: str, size: float, width: float) -> list[str]:
    """تقسيم النص إلى أسطر حسب العرض (بالترتيب المنطقي)، ثم تشكيل كل سطر."""
    import arabic_reshaper
    from reportlab.pdfbase.pdfmetrics import stringWidth

    lines, current = [], ""
    for word in str(text).split():
        candidate = f"{current} {word}".strip()
        if stringWidth(arabic_reshaper.reshape(candidate), font, size) <= width or not current:
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return [shape_arabic(line) for line in lines] or [""]


def export_pdf(data: dict) -> bytes:
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    _register_fonts()
    page_width = A4[0] - 4 * cm
    body = ParagraphStyle("body", fontName="Amiri", fontSize=11, leading=17, alignment=TA_RIGHT)
    heading = ParagraphStyle("heading", parent=body, fontName="Amiri-Bold", fontSize=15, leading=24,
                             textColor=colors.HexColor("#1F5F4A"), spaceBefore=10, spaceAfter=4)
    title_style = ParagraphStyle("title", parent=heading, fontSize=20, leading=30, alignment=1)
    cell_style = ParagraphStyle("cell", parent=body, fontSize=9.5, leading=13)

    def para(text, style=body, width=page_width):
        lines = wrap_rtl(text, style.fontName, style.fontSize, width)
        return Paragraph("<br/>".join(escape(line) for line in lines), style)

    def table(headers, rows, widths):
        """جدول من اليمين لليسار: نعكس ترتيب الأعمدة."""
        headers, widths = headers[::-1], widths[::-1]
        content = [[para(h, cell_style, w - 14) for h, w in zip(headers, widths)]]
        for row in rows:
            content.append([para(str(v), cell_style, w - 14) for v, w in zip(row[::-1], widths)])
        t = Table(content, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCEFE6")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9DB8AC")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        return t

    story = [para(data["title"], title_style),
             para(f"تاريخ الإنشاء: {data['generated_at']}", body)]
    for caveat in data["caveats"][:1]:
        if caveat.startswith("⚠"):
            story.append(para(caveat.replace("⚠", "تنبيه:"), ParagraphStyle(
                "warn", parent=body, textColor=colors.HexColor("#B71C1C"))))

    story.append(para("الملخص الإحصائي", heading))
    for line in summary_lines(data):
        story.append(para("• " + line))

    if data["n_reviews"]:
        dist = data["distribution"]
        drawing = Drawing(page_width, 170)
        chart = VerticalBarChart()
        chart.x, chart.y, chart.width, chart.height = 40, 30, page_width - 80, 120
        chart.data = [dist["count"].tolist()]
        chart.categoryAxis.categoryNames = [shape_arabic(x) for x in dist["label"]]
        chart.categoryAxis.labels.fontName = "Amiri"
        chart.valueAxis.labels.fontName = "Amiri"
        chart.valueAxis.valueMin = 0
        chart.bars[0].fillColor = colors.HexColor("#3A8F6E")
        for i, s in enumerate(dist["sentiment"]):
            chart.bars[(0, i)].fillColor = colors.HexColor(config.SENTIMENT_COLORS[s])
        drawing.add(chart)
        story.append(drawing)

    if not data["aspects"].empty:
        story.append(para("جوانب الرأي", heading))
        rows = [(r["aspect_ar"], r["n"], fraction_text(int(r["positive"]), int(r["n"])),
                 fraction_text(int(r["negative"]), int(r["n"])))
                for r in data["aspects"].to_dict("records")]
        story.append(table(["الجانب", "الإشارات", "إيجابي", "سلبي"], rows,
                           [page_width * w for w in (0.3, 0.14, 0.28, 0.28)]))

    if not data["dishes"].empty:
        story.append(para("الوجبات الأكثر ذكرًا (مؤشر تفضيل وليس مبيعات)", heading))
        rows = [(r["dish"], r["mentions"], fraction_text(int(r["positive"]), int(r["mentions"])),
                 fraction_text(int(r["negative"]), int(r["mentions"])))
                for r in data["dishes"].head(10).to_dict("records")]
        story.append(table(["الوجبة", "تعليقات ذكرتها", "إيجابي", "سلبي"], rows,
                           [page_width * w for w in (0.3, 0.14, 0.28, 0.28)]))

    for key, label in (("complaints", "أمثلة على الشكاوى"), ("praises", "أمثلة على الإشادات")):
        frame = data[key]
        if frame is not None and not frame.empty:
            story.append(para(label, heading))
            rows = [(config.ASPECT_AR.get(r["aspect"], r["aspect"]), r["evidence"])
                    for r in frame.head(6).to_dict("records")]
            story.append(table(["الجانب", "العبارة الداعمة"], rows,
                               [page_width * 0.25, page_width * 0.75]))

    story.append(para("التوصيات", heading))
    for rec in data["recommendations"]:
        story.append(para("• " + rec))
    story.append(para("تنبيهات منهجية", heading))
    for caveat in data["caveats"]:
        story.append(para("• " + caveat.replace("⚠ ", "")))
    story.append(Spacer(1, 12))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=2 * cm, leftMargin=2 * cm,
                            topMargin=1.8 * cm, bottomMargin=1.8 * cm, title="Restaurant Sentiment Report")
    doc.build(story)
    return buffer.getvalue()


def export_csv(data: dict) -> bytes:
    if not data["n_reviews"]:
        return "".encode("utf-8-sig")
    return data["reviews"].to_csv(index=False).encode("utf-8-sig")
